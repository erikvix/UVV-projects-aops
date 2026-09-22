#!/usr/bin/env python3
"""Gera, a partir de dados/coletas.csv, tudo o que e publicado na AOP3:

  * aop3/site/js/dados.js  -> dados que alimentam o site
  * aop3/planilhas/*.csv   -> uma planilha por consulta (item II.d)
  * aop3/planilhas/precos-combustiveis-vila-velha.xlsx
        -> planilha unica com as 4 consultas, as coletas e os 2 graficos do item II.e

Uso:  python3 aop3/scripts/gerar_site_e_planilhas.py
"""
from __future__ import annotations

import csv
import json
from collections import OrderedDict, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[2]
ENTRADA = RAIZ / "dados" / "coletas.csv"
FONTE_TXT = RAIZ / "dados" / "fonte.txt"
SITE_JS = RAIZ / "aop3" / "site" / "js" / "dados.js"
PLANILHAS = RAIZ / "aop3" / "planilhas"
XLSX = PLANILHAS / "precos-combustiveis-vila-velha.xlsx"

CIDADE, UF = "Vila Velha", "ES"

# Enquanto True, o site exibe o aviso de que os precos ainda sao provisorios.
# Troque para False depois de publicar a coleta real (ou a importacao da ANP).
DADOS_DE_DEMONSTRACAO = False

FONTE_DEMONSTRACAO = (
    "Conjunto de demonstracao gerado por aop3/scripts/gerar_dados_exemplo.py, "
    "com 5 postos ficticios distribuidos em bairros reais de Vila Velha/ES. "
    "Serve apenas para validar o site: substitua dados/coletas.csv pela coleta de campo "
    "do grupo ou pela Serie Historica de Precos de Combustiveis da ANP "
    "(aop3/scripts/importar_anp.py) antes da divulgacao."
)
FONTE_REAL = (
    "Coleta realizada pelo grupo em postos de Vila Velha/ES para o Projeto de Extensao da "
    "disciplina Arquitetura de Dados Relacionais I (UVV). Os dados sao armazenados no banco "
    "relacional construido na AOP2 e exportados para esta pagina."
)

CABECALHO = PatternFill("solid", fgColor="E8EEF7")
NEGRITO = Font(bold=True)


# --------------------------------------------------------------------------
# leitura
# --------------------------------------------------------------------------
def ler_coletas() -> list[dict]:
    with ENTRADA.open(encoding="utf-8", newline="") as fp:
        coletas = [
            {
                "posto": l["posto"],
                "bairro": l["bairro"],
                "endereco": l["endereco"],
                "bandeira": l["bandeira"],
                "combustivel": l["combustivel"],
                "data": l["data_coleta"],
                "preco": round(float(l["preco"]), 3),
            }
            for l in csv.DictReader(fp)
        ]
    if not coletas:
        raise SystemExit(f"{ENTRADA} esta vazio")
    return sorted(coletas, key=lambda c: (c["posto"], c["combustivel"], c["data"]))


def descricao_da_fonte() -> str:
    """dados/fonte.txt e escrito por importar_anp.py e descreve exatamente de
    onde vieram as coletas; sem ele, cai nos textos padrao acima."""
    if FONTE_TXT.exists():
        texto = FONTE_TXT.read_text(encoding="utf-8").strip()
        if texto:
            return texto
    return FONTE_DEMONSTRACAO if DADOS_DE_DEMONSTRACAO else FONTE_REAL


def ordenados(valores) -> list[str]:
    return sorted(set(valores), key=lambda v: v.lower())


def br_data(iso: str) -> str:
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%d/%m/%Y")


# --------------------------------------------------------------------------
# as quatro consultas do item II.d (mesmas regras do SQL e do site)
# --------------------------------------------------------------------------
def consulta_1(coletas):
    """Menor e maior preco de cada combustivel. Empate: coleta mais recente."""
    linhas = []
    for combustivel in ordenados(c["combustivel"] for c in coletas):
        # da mais recente para a mais antiga: min()/max() sao estaveis, entao o
        # empate de preco fica com a coleta mais recente (mesma regra do SQL)
        lista = sorted((c for c in coletas if c["combustivel"] == combustivel),
                       key=lambda c: c["data"], reverse=True)
        menor = min(lista, key=lambda c: c["preco"])
        maior = max(lista, key=lambda c: c["preco"])
        for tipo, c in (("Menor preco", menor), ("Maior preco", maior)):
            linhas.append([combustivel, tipo, c["posto"], c["endereco"], c["bairro"],
                           c["preco"], br_data(c["data"])])
    return linhas


def consulta_2(coletas):
    """Preco medio e quantidade de amostras por posto e combustivel."""
    grupos = defaultdict(list)
    for c in coletas:
        grupos[(c["posto"], c["combustivel"])].append(c)
    linhas = []
    for (posto, combustivel), lista in grupos.items():
        media = sum(c["preco"] for c in lista) / len(lista)
        linhas.append([posto, lista[0]["bairro"], combustivel, round(media, 3), len(lista)])
    return sorted(linhas, key=lambda l: (l[2].lower(), l[3]))


def consulta_3(coletas):
    """Preco mais recente de cada combustivel em cada posto."""
    ultimas = {}
    for c in coletas:
        chave = (c["posto"], c["combustivel"])
        if chave not in ultimas or c["data"] > ultimas[chave]["data"]:
            ultimas[chave] = c
    linhas = [[c["posto"], c["bairro"], c["combustivel"], c["preco"], br_data(c["data"])]
              for c in ultimas.values()]
    return sorted(linhas, key=lambda l: (l[2].lower(), l[3]))


def consulta_4(coletas):
    """Evolucao do preco ao longo do tempo, por posto e combustivel."""
    linhas = [[c["posto"], c["bairro"], c["combustivel"], c["preco"], br_data(c["data"])]
              for c in sorted(coletas, key=lambda c: (c["posto"].lower(), c["combustivel"].lower(), c["data"]))]
    return linhas


def media_por_data(coletas, posto=None):
    """{combustivel: [(data, media), ...]} - base dos graficos do item II.e."""
    base = [c for c in coletas if posto is None or c["posto"] == posto]
    soma = defaultdict(lambda: [0.0, 0])
    for c in base:
        item = soma[(c["combustivel"], c["data"])]
        item[0] += c["preco"]
        item[1] += 1
    resultado = OrderedDict()
    for combustivel in ordenados(c["combustivel"] for c in base):
        pontos = [(data, round(v[0] / v[1], 3))
                  for (comb, data), v in sorted(soma.items(), key=lambda kv: kv[0][1])
                  if comb == combustivel]
        resultado[combustivel] = pontos
    return resultado


# --------------------------------------------------------------------------
# saidas
# --------------------------------------------------------------------------
COLUNAS = {
    1: ["Combustivel", "Tipo", "Posto", "Endereco", "Bairro", "Preco (R$)", "Data da coleta"],
    2: ["Posto", "Bairro", "Combustivel", "Preco medio (R$)", "Qtd. de amostras"],
    3: ["Posto", "Bairro", "Combustivel", "Preco (R$)", "Data da coleta"],
    4: ["Posto", "Bairro", "Combustivel", "Preco (R$)", "Data da coleta"],
}
TITULOS = {
    1: "Consulta I - menor e maior preco de cada combustivel",
    2: "Consulta II - preco medio e quantidade de amostras por posto",
    3: "Consulta III - preco mais recente de cada combustivel em cada posto",
    4: "Consulta IV - evolucao do preco ao longo do tempo",
}


def gravar_dados_js(coletas) -> None:
    datas = sorted(c["data"] for c in coletas)
    payload = {
        "cidade": CIDADE,
        "uf": UF,
        "demo": DADOS_DE_DEMONSTRACAO,
        "fonte": descricao_da_fonte(),
        "gerado_em": date.today().isoformat(),
        "periodo": {"inicio": datas[0], "fim": datas[-1]},
        "coletas": coletas,
    }
    SITE_JS.parent.mkdir(parents=True, exist_ok=True)
    SITE_JS.write_text(
        "/* Arquivo gerado por aop3/scripts/gerar_site_e_planilhas.py - nao edite a mao. */\n"
        "window.DADOS = " + json.dumps(payload, ensure_ascii=False, indent=1) + ";\n",
        encoding="utf-8",
    )


def gravar_csvs(resultados) -> None:
    PLANILHAS.mkdir(parents=True, exist_ok=True)
    for numero, linhas in resultados.items():
        caminho = PLANILHAS / f"consulta-{numero}.csv"
        with caminho.open("w", encoding="utf-8-sig", newline="") as fp:
            escritor = csv.writer(fp, delimiter=";")
            escritor.writerow(COLUNAS[numero])
            escritor.writerows(linhas)


def ajustar(planilha, larguras) -> None:
    for i, largura in enumerate(larguras, start=1):
        planilha.column_dimensions[get_column_letter(i)].width = largura


def aba_tabela(wb, titulo_aba, titulo, colunas, linhas, larguras) -> None:
    ws = wb.create_sheet(titulo_aba)
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(colunas)
    for celula in ws[3]:
        celula.font = NEGRITO
        celula.fill = CABECALHO
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for linha in linhas:
        ws.append(linha)
    for coluna, titulo_col in enumerate(colunas, start=1):
        if "R$" in titulo_col:
            for celula in ws.iter_rows(min_row=4, min_col=coluna, max_col=coluna):
                celula[0].number_format = '"R$" #,##0.000'
    ajustar(ws, larguras)
    ws.freeze_panes = "A4"


def aba_grafico(wb, titulo_aba, titulo, series_por_bloco) -> None:
    """series_por_bloco: [(subtitulo, {combustivel: [(data, media)]})]"""
    ws = wb.create_sheet(titulo_aba)
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    linha = 3
    for subtitulo, series in series_por_bloco:
        if not any(series.values()):
            continue
        combustiveis = [c for c, pontos in series.items() if pontos]
        datas = sorted({d for pontos in series.values() for d, _ in pontos})
        ws.cell(row=linha, column=1, value=subtitulo).font = Font(bold=True, size=11)
        linha += 1
        inicio = linha
        ws.cell(row=linha, column=1, value="Data da coleta").font = NEGRITO
        ws.cell(row=linha, column=1).fill = CABECALHO
        for j, combustivel in enumerate(combustiveis, start=2):
            celula = ws.cell(row=linha, column=j, value=combustivel)
            celula.font = NEGRITO
            celula.fill = CABECALHO
        for data_iso in datas:
            linha += 1
            ws.cell(row=linha, column=1, value=br_data(data_iso))
            for j, combustivel in enumerate(combustiveis, start=2):
                valor = dict(series[combustivel]).get(data_iso)
                celula = ws.cell(row=linha, column=j, value=valor)
                celula.number_format = '"R$" #,##0.000'

        gr = LineChart()
        gr.title = subtitulo
        gr.y_axis.title = "Preco medio (R$)"
        gr.x_axis.title = "Data da coleta"
        gr.height, gr.width = 8.5, 19
        gr.add_data(Reference(ws, min_col=2, max_col=1 + len(combustiveis),
                              min_row=inicio, max_row=linha), titles_from_data=True)
        gr.set_categories(Reference(ws, min_col=1, min_row=inicio + 1, max_row=linha))
        ws.add_chart(gr, f"H{inicio}")
        linha += 20  # espaco para o grafico antes do proximo bloco
    ajustar(ws, [18] + [20] * 6)


def gravar_xlsx(coletas, resultados) -> None:
    wb = Workbook()
    leia = wb.active
    leia.title = "Leia-me"
    postos = ordenados(c["posto"] for c in coletas)
    bairros = ordenados(c["bairro"] for c in coletas)
    datas = sorted(c["data"] for c in coletas)
    info = [
        ("Precos de combustiveis - " + CIDADE + "/" + UF, ""),
        ("", ""),
        ("Projeto", "Projeto de Extensao - Arquitetura de Dados Relacionais I (UVV)"),
        ("Entrega", "AOP3 - divulgacao do trabalho para a comunidade"),
        ("Gerado em", date.today().strftime("%d/%m/%Y")),
        ("Periodo das coletas", f"{br_data(datas[0])} a {br_data(datas[-1])}"),
        ("Coletas", len(coletas)),
        ("Postos", f"{len(postos)} ({', '.join(postos)})"),
        ("Bairros", f"{len(bairros)} ({', '.join(bairros)})"),
        ("Combustiveis", ", ".join(ordenados(c["combustivel"] for c in coletas))),
        ("Fonte dos dados", descricao_da_fonte()),
        ("", ""),
        ("Abas desta planilha", "Consulta I a Consulta IV: as consultas exigidas no item II.d"),
        ("", "Grafico - cidade / Grafico - postos: os graficos do item II.e"),
        ("", "Coletas: todas as coletas, uma por linha"),
    ]
    for rotulo, valor in info:
        leia.append([rotulo, valor])
    leia["A1"].font = Font(bold=True, size=14)
    for linha in leia.iter_rows(min_row=3, max_col=1):
        linha[0].font = NEGRITO
    for linha in leia.iter_rows(min_row=1, min_col=2, max_col=2):
        linha[0].alignment = Alignment(wrap_text=True, vertical="top")
    ajustar(leia, [22, 95])

    aba_tabela(wb, "Consulta I", TITULOS[1], COLUNAS[1], resultados[1], [22, 14, 24, 42, 24, 14, 16])
    aba_tabela(wb, "Consulta II", TITULOS[2], COLUNAS[2], resultados[2], [24, 24, 22, 18, 18])
    aba_tabela(wb, "Consulta III", TITULOS[3], COLUNAS[3], resultados[3], [24, 24, 22, 14, 16])
    aba_tabela(wb, "Consulta IV", TITULOS[4], COLUNAS[4], resultados[4], [24, 24, 22, 14, 16])

    aba_grafico(wb, "Grafico - cidade",
                "Evolucao do preco medio de cada combustivel em " + CIDADE,
                [("Media da cidade - todos os postos", media_por_data(coletas))])
    aba_grafico(wb, "Grafico - postos",
                "Evolucao do preco medio de cada combustivel em cada posto",
                [(posto, media_por_data(coletas, posto)) for posto in ordenados(c["posto"] for c in coletas)])

    aba_tabela(wb, "Coletas", "Todas as coletas registradas",
               ["Posto", "Bairro", "Endereco", "Bandeira", "Combustivel", "Data da coleta", "Preco (R$)"],
               [[c["posto"], c["bairro"], c["endereco"], c["bandeira"], c["combustivel"],
                 br_data(c["data"]), c["preco"]] for c in coletas],
               [24, 24, 42, 18, 22, 16, 14])

    PLANILHAS.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)


def main() -> None:
    coletas = ler_coletas()
    resultados = {
        1: consulta_1(coletas),
        2: consulta_2(coletas),
        3: consulta_3(coletas),
        4: consulta_4(coletas),
    }
    gravar_dados_js(coletas)
    gravar_csvs(resultados)
    gravar_xlsx(coletas, resultados)
    print(f"{len(coletas)} coletas processadas")
    print(f"  site     -> {SITE_JS.relative_to(RAIZ)}")
    print(f"  planilhas-> {PLANILHAS.relative_to(RAIZ)}/ (4 CSVs + {XLSX.name})")


if __name__ == "__main__":
    main()
